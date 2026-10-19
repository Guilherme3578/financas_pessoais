from datetime import date
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.repositories.relatorio import RelatorioRepository


class RelatorioService:
    def __init__(self, db: Session):
        self.repository = RelatorioRepository(db)

    def gerar_excel(
        self,
        data_inicio: date,
        data_fim: date,
        tipo: str | None = None,
        status: str | None = None,
    ) -> BytesIO:
        if data_inicio > data_fim:
            raise HTTPException(
                status_code=422,
                detail="A data inicial não pode ser posterior à data final.",
            )

        registros = self.repository.listar_lancamentos(
            data_inicio=data_inicio,
            data_fim=data_fim,
            tipo=tipo,
            status=status,
        )

        if len(registros) > 10000:
            raise HTTPException(
                status_code=422,
                detail=(
                    "O período possui mais de 10.000 lançamentos. "
                    "Selecione um período menor."
                ),
            )

        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Lançamentos"

        cabecalhos = [
            "ID",
            "Descrição",
            "Tipo",
            "Categoria",
            "Valor",
            "Competência",
            "Pagamento / recebimento previsto",
            "Pagamento / recebimento realizado",
            "Status",
            "Forma de pagamento",
            "Observação",
        ]

        planilha.append(cabecalhos)
        planilha.freeze_panes = "A2"
        planilha.sheet_view.showGridLines = False

        formas = {
            "pix": "Pix",
            "debito": "Débito",
            "dinheiro": "Dinheiro",
            "transferencia": "Transferência",
            "boleto": "Boleto",
        }

        for lancamento, categoria_nome in registros:
            if lancamento.status == "efetivado":
                situacao = (
                    "Recebido" if lancamento.tipo == "receita" else "Pago"
                )
            elif lancamento.status == "cancelado":
                situacao = "Cancelado"
            else:
                situacao = "Pendente"

            planilha.append(
                [
                    lancamento.id,
                    lancamento.descricao,
                    "Receita" if lancamento.tipo == "receita" else "Despesa",
                    categoria_nome,
                    lancamento.valor,
                    lancamento.data_competencia,
                    lancamento.data_vencimento,
                    lancamento.data_efetivacao,
                    situacao,
                    formas.get(
                        lancamento.forma_pagamento,
                        lancamento.forma_pagamento,
                    ),
                    lancamento.observacao,
                ]
            )

            numero_linha = planilha.max_row

            for coluna in (2, 3, 4, 9, 10, 11):
                celula = planilha.cell(numero_linha, coluna)

                if celula.value is not None:
                    celula.data_type = "s"

            planilha.cell(numero_linha, 5).number_format = (
                '"R$" #,##0.00'
            )

            for coluna in (6, 7, 8):
                planilha.cell(numero_linha, coluna).number_format = (
                    "dd/mm/yyyy"
                )

            for celula in planilha[numero_linha]:
                celula.alignment = Alignment(vertical="center")

                if numero_linha % 2 == 0:
                    celula.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F1F6F3",
                    )

        for celula in planilha[1]:
            celula.font = Font(
                name="Calibri",
                bold=True,
                color="FFFFFF",
            )
            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="174F3E",
            )
            celula.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

        planilha.row_dimensions[1].height = 32

        larguras = [10, 38, 14, 26, 20, 18, 28, 28, 16, 24, 48]

        for indice, largura in enumerate(larguras, start=1):
            letra = get_column_letter(indice)
            planilha.column_dimensions[letra].width = largura

        planilha.auto_filter.ref = f"A1:K{planilha.max_row}"

        informacoes = workbook.create_sheet("Informações")
        informacoes.append(["Exportação de lançamentos"])
        informacoes.append(["Início da competência", data_inicio])
        informacoes.append(["Fim da competência", data_fim])
        informacoes.append(["Tipo", tipo or "Receitas e despesas"])
        informacoes.append(["Status", status or "Todos, incluindo cancelados"])
        informacoes.append(["Quantidade de lançamentos", len(registros)])
        informacoes.append(
            [
                "Critério",
                "Período filtrado pela data de competência.",
            ]
        )

        informacoes["B2"].number_format = "dd/mm/yyyy"
        informacoes["B3"].number_format = "dd/mm/yyyy"
        informacoes.column_dimensions["A"].width = 30
        informacoes.column_dimensions["B"].width = 58
        informacoes["A1"].font = Font(bold=True, color="174F3E", size=14)

        arquivo = BytesIO()
        workbook.save(arquivo)
        workbook.close()
        arquivo.seek(0)

        return arquivo